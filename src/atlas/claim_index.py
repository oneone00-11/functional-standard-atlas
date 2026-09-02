"""Where else is this claim stated?

A correction applied in one place does not travel. The clinical-yield sentence
in the companion manuscript was corrected in Results and left standing,
unqualified, in the Abstract, the Introduction and the Discussion -- three
restatements that survived because nothing listed them. This builds that list.

It is deliberately not the analysis-claims manifest. config/analysis_claims.yaml
asks whether the analysis a sentence describes exists; this asks where else the
same claim is made, so that a fix can be carried to all of them.

    python -m atlas.claim_index --report <manuscript.docx>
    python -m atlas.claim_index --report <manuscript.docx> --claim clinical-yield
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml

REPO = Path(__file__).resolve().parents[2]
CONFIG = REPO / "config" / "claim_restatements.yaml"


def load(config: Path = CONFIG) -> dict:
    return yaml.safe_load(config.read_text())["manuscripts"]


def sections(doc) -> dict[int, str]:
    """paragraph index -> the heading it sits under."""
    out, cur = {}, "(front matter)"
    for i, p in enumerate(doc.paragraphs):
        if p.style.name.startswith("Heading") and p.text.strip():
            cur = p.text.strip()
        out[i] = cur
    return out


def index(docx_path: Path, claims: list[dict]) -> dict[str, list[dict]]:
    import docx

    d = docx.Document(str(docx_path))
    sec = sections(d)
    units = [(f"P{i}", sec[i], p.text) for i, p in enumerate(d.paragraphs) if p.text.strip()]
    for ti, t in enumerate(d.tables):
        for ri, row in enumerate(t.rows):
            units.append((f"T{ti + 1}r{ri}", "(table)",
                          " | ".join(c.text.strip() for c in row.cells)))

    found: dict[str, list[dict]] = {}
    for c in claims:
        pat = re.compile(c["pattern"], re.I)
        hits = []
        for uid, section, text in units:
            if re.match(r"^\s*\d{1,2}\.\s+\S+.*(doi|https?://)", text):
                continue                       # reference-list entry
            for s in re.split(r"(?<=[.;])\s+(?=[A-Z(0-9])", text):
                if pat.search(s):
                    hits.append({"where": uid, "section": section,
                                 "sentence": re.sub(r"\s+", " ", s).strip()})
        found[c["id"]] = hits
    return found


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", required=True, metavar="DOCX")
    ap.add_argument("--claim", help="only this claim id")
    a = ap.parse_args()

    path = Path(a.report)
    name = path.name
    specs = load()
    spec = next((v for v in specs.values() if v["file"] == name), None)
    if spec is None:
        sys.exit(f"[claims] {name} is not listed in {CONFIG.relative_to(REPO)}; "
                 "add it before relying on this report")
    claims = [c for c in spec["claims"] if not a.claim or c["id"] == a.claim]
    if not claims:
        sys.exit(f"[claims] no claim with id {a.claim!r}")

    idx = index(path, claims)
    print(f"[claims] {name}")
    for c in claims:
        hits = idx[c["id"]]
        secs = sorted({h["section"] for h in hits})
        print(f"\n  {c['id']}  —  {c['says']}")
        print(f"    {len(hits)} restatement(s) across {len(secs)} section(s): {', '.join(secs)}")
        for h in hits:
            print(f"      [{h['where']}] ({h['section']}) {h['sentence'][:150]}")
    print("\n[claims] fixing any one of these means checking every location listed above.")


if __name__ == "__main__":
    main()
