"""Build the related-manuscript deliverable from the companion manuscript.

The delivered copy sat three months behind its source, still carrying a claim
the companion manuscript had retracted, because it was exported by hand and
nothing tied the two together. This does the mechanical part: copy, prepend the
status page, strip font declarations a converter cannot embed, and check the
markers that say whether the export is current.

The status page names no venue and no submission state. Those change with every
submission and belong in the cover letter, which is rewritten each time; the
relationship between the two manuscripts does not change.

Usage (PYTHONPATH=src):
    python -m atlas.build_related_manuscript --source ~/Desktop/<companion>.docx \
                                             --out <staging>/Supplemental_Related_Manuscript_1.docx
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from atlas.supplement import _drop_restricted_fonts

STATUS = ("Status: Companion manuscript by the same author. It shares source data with "
          "this submission but no results section, figure or table.")

# What a current export must and must not contain. Kept here rather than in the
# package manifest because the retracted phrasings are the companion paper's own
# text, and the manifest is published.
MUST_CONTAIN = ["neither changes the evidence strength", "Kish effective sample size",
                "Forty-eight automated tests", "ascertainment rather than circularity"]
MUST_NOT_CONTAIN = ["more actionable calls without loss of accuracy",
                    "Ranking is saturated but calibration is not"]


def build(source: Path, out: Path, author: str = "Ningyi Zhang") -> dict:
    from docx import Document
    from docx.enum.text import WD_BREAK

    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, out)          # the source is never edited

    d = Document(str(out))
    anchor = d.paragraphs[0]._p           # the manuscript's own title
    body = d.paragraphs[1].style

    def insert(text: str = "", bold: bool = False, page_break: bool = False) -> None:
        p = d.add_paragraph(text, style=body)
        if bold and p.runs:
            p.runs[0].bold = True
        if page_break:
            p.add_run().add_break(WD_BREAK.PAGE)
        anchor.addprevious(p._p)

    insert("Related manuscript", bold=True)
    insert()
    insert(f"Author: {author}")
    insert()
    insert(STATUS)
    insert(page_break=True)
    d.save(str(out))

    fixed = _drop_restricted_fonts(out)
    return {"fonts_fixed": fixed}


def check(path: Path) -> dict:
    from atlas.submission_package import read_docx

    t = read_docx(path)
    return {"missing": [m for m in MUST_CONTAIN if m not in t],
            "should_be_gone": [m for m in MUST_NOT_CONTAIN if m in t],
            "status_present": STATUS in t,
            "words": len(t.split())}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    src, out = Path(a.source).expanduser(), Path(a.out).expanduser()
    if not src.exists():
        sys.exit(f"[related] source not found: {src}")

    r = build(src, out)
    if r["fonts_fixed"]:
        print(f"[related] removed a restricted font declaration from: "
              f"{', '.join(r['fonts_fixed'])}")
    c = check(out)
    print(f"[related] wrote {out} ({c['words']:,} words)")
    print(f"[related] status page present: {c['status_present']}")
    for m in c["missing"]:
        print(f"  MISSING from the export: {m!r}")
    for m in c["should_be_gone"]:
        print(f"  STILL PRESENT (retracted): {m!r}")
    ok = c["status_present"] and not c["missing"] and not c["should_be_gone"]
    print(f"[related] {'all markers correct' if ok else 'markers FAILED'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
