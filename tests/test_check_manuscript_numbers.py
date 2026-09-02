"""Guardrails for the manuscript number checker.

The checker exists because the manuscript claims every reported value is read
programmatically from pipeline output, and nothing enforced it. These tests pin
the two behaviours that make it worth having:

  * a token with no counterpart anywhere in results/ must be reported, never
    silently skipped -- the whitelist is the only way past, and it is tracked;
  * a match that carries no information must not be counted as verification.
    A two-decimal number in a dense range matches hundreds of unrelated pipeline
    values; calling that "verified" is precisely how the complementarity range
    survived six rounds of checking.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import numpy as np
import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from atlas.check_manuscript_numbers import (  # noqa: E402
    CITATION, IDENTIFIER, TOKEN, WEAK_MATCH_MIN, WHITELIST, _tolerance,
    load_whitelist, pipeline_values,
)

RESULTS = REPO / "results"


def test_whitelist_is_tracked_and_every_entry_carries_a_reason():
    assert WHITELIST.exists(), "the whitelist must be committed, not implicit"
    _, _, spec = load_whitelist()
    assert spec, "whitelist is empty; it should at least document why"
    for name, group in spec.items():
        assert group.get("reason", "").strip(), f"{name} has no reason"
        assert group.get("values") or group.get("pattern"), f"{name} matches nothing"


def test_tolerance_follows_printed_precision():
    assert _tolerance("0.12") == pytest.approx(0.005, abs=1e-9)
    assert _tolerance("0.1234") == pytest.approx(0.00005, abs=1e-10)
    assert _tolerance("17") == pytest.approx(0.5, abs=1e-9)


def test_identifiers_and_citations_are_not_read_as_measurements():
    """A DOI and a citation number are not values the pipeline must explain."""
    text = ("available at https://doi.org/10.5281/zenodo.21828448 [61] and "
            "urn:mavedb:00000097-0-2, see [22, 23, 24]")
    stripped = CITATION.sub(" ", IDENTIFIER.sub(" ", text))
    assert TOKEN.search(stripped) is None, TOKEN.search(stripped).group(0)


def test_thousands_separated_tokens_are_read():
    m = TOKEN.search("across 64,178 variants")
    assert m and m.group(1) == "64,178"


def test_row_counts_enter_the_pool():
    """"2,452 of 2,803 score sets" are real pipeline quantities that appear
    nowhere inside the files themselves."""
    if not RESULTS.is_dir() or not (RESULTS / "mavedb_survey_v1.tsv").exists():
        pytest.skip("results/ not built in this checkout")
    pool = pipeline_values()
    for n in (2452.0, 2803.0):
        assert np.any(np.abs(pool - n) <= 0.5), f"{n:.0f} missing from the pool"


def test_generated_prose_is_excluded_from_the_pool():
    """Supplemental_Note.md is produced by atlas.supplement; a literal hardcoded
    in that generator would otherwise validate itself through it."""
    src = (REPO / "src" / "atlas" / "check_manuscript_numbers.py").read_text()
    fn = src.split("def pipeline_values")[1].split("\ndef ")[0]
    assert '".md"' not in fn, "generated prose must stay out of the pool"


def test_weak_matches_are_a_separate_class():
    """The limit of this tool, pinned so it is not forgotten: a low-precision
    number in a dense range matches by coincidence, so the checker alone could
    not have caught the complementarity range. That binding is what the
    analysis-claims manifest provides."""
    if not RESULTS.is_dir() or not any(RESULTS.iterdir()):
        pytest.skip("results/ not built in this checkout")
    pool = pipeline_values()
    n = int(np.sum(np.abs(pool - 0.12) <= 0.005))
    assert n >= WEAK_MATCH_MIN, (
        "0.12 no longer matches many pipeline values; the weak-match threshold "
        "may need revisiting")


# --- the cardinality gate ---------------------------------------------------
# The complementarity sentence is the reason this exists. No value check catches
# it: every number it prints corresponds to some cell of the matrix it cites,
# whether the check runs against the whole pipeline or against that one file.
# What was wrong was the COUNT -- "the nineteen predictors" citing a nine-model
# matrix. That is a claim about the output's shape, and only a shape check sees
# it. Verified against the real pre-correction manuscript at the time this was
# written (1 cardinality mismatch, 0 after the fix); reproduced here on a
# synthetic document so the suite stays self-contained.

def _doc(tmp_path, text):
    import docx
    d = docx.Document()
    d.add_paragraph("padding")          # so the sentence lands in P1
    d.add_paragraph(text)
    p = tmp_path / "m.docx"
    d.save(p)
    return p


def _claims(tmp_path, count):
    p = tmp_path / "claims.yaml"
    p.write_text(
        "claims:\n"
        "  - para: P1\n"
        "    claim: two axes\n"
        "    script: src/atlas/ensemble.py\n"
        "    output: results/model_correlation_coding_or_utr_v1.tsv\n"
        "    status: verified\n"
        f"    cardinality: {{count: {count}, noun: predictors}}\n")
    return p


def test_cardinality_mismatch_is_caught(tmp_path):
    from atlas.check_manuscript_numbers import classify
    doc = _doc(tmp_path, "The matrix shows that the nineteen predictors span two axes.")
    r = classify(doc, claims=_claims(tmp_path, 9))
    assert len(r["cardinality_mismatch"]) == 1, r["cardinality_mismatch"]
    m = r["cardinality_mismatch"][0]
    assert m["asserted"] == 19 and m["recorded"] == 9


def test_correct_cardinality_passes(tmp_path):
    from atlas.check_manuscript_numbers import classify
    doc = _doc(tmp_path, "The matrix shows that the nine broad-scope and "
                         "splice-aware predictors span two axes.")
    r = classify(doc, claims=_claims(tmp_path, 9))
    assert r["cardinality_mismatch"] == []


def test_value_checks_alone_would_not_have_caught_it(tmp_path):
    """Pinned so the cardinality gate is never removed as redundant: the printed
    numbers in that sentence all correspond to real cells, scoped or not."""
    import numpy as np
    from atlas.check_manuscript_numbers import claim_scopes, scoped_values
    scopes = claim_scopes()
    if "P71" not in scopes or not (REPO / scopes["P71"][0]).exists():
        pytest.skip("results/ not built in this checkout")
    pool = scoped_values(scopes["P71"])
    assert np.any(np.abs(pool - 0.12) <= 0.005), (
        "0.12 no longer corresponds to a cell of the cited matrix; the note in "
        "this test about why a value check is insufficient may need revisiting")


# --- declared counts about the pipeline itself -------------------------------
# The guardrail count is the case that proves pool membership is the wrong test.
# 102 (the count before this work) and 118 (the count during it) each match
# exactly one unrelated pipeline value, so a stale count passes as VERIFIED.
# It is a declared fact, so it is compared against the declared value.

def test_stale_declared_count_is_caught(tmp_path):
    from atlas.check_manuscript_numbers import classify
    doc = _doc(tmp_path, "the pipeline is covered by 118 guardrail tests that run from a clean extract")
    r = classify(doc)
    assert len(r["stale_counts"]) == 1, r["stale_counts"]
    assert r["stale_counts"][0]["stated"] == 118


def test_current_declared_count_passes(tmp_path):
    """The key checked is the ARCHIVE-measured passing count, not the local
    collected count: the manuscript says the tests run "from a clean extract",
    and a development tree is not one."""
    import json
    from atlas.check_manuscript_numbers import DECLARED_COUNTS, FACTS, classify
    key = next(iter(DECLARED_COUNTS.values()))
    n = json.loads(FACTS.read_text())[key]
    doc = _doc(tmp_path, f"the pipeline is covered by {n} guardrail tests")
    assert classify(doc)["stale_counts"] == []


def test_a_stale_count_would_otherwise_pass_as_verified(tmp_path):
    """Pinned so this check is never dropped as redundant with the value pool."""
    import numpy as np
    from atlas.check_manuscript_numbers import pipeline_values
    if not RESULTS.is_dir() or not any(RESULTS.iterdir()):
        pytest.skip("results/ not built in this checkout")
    pool = pipeline_values()
    for stale in (102, 118):
        assert np.any(np.abs(pool - stale) <= 0.5), (
            f"{stale} no longer coincides with a pipeline value; the rationale "
            "recorded here may need revisiting")
