"""Guardrails for the claim-restatement index.

A correction applied in one place does not travel. The clinical-yield sentence
in the companion manuscript was corrected in Results and left standing,
unqualified, in the Abstract, the Introduction and the Discussion. Nothing
listed those other locations, so nothing carried the fix to them.

config/claim_restatements.yaml is that list. These tests keep it usable: every
claim carries a plain-language statement of what it asserts, and every pattern
compiles and is specific enough to be worth running.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from atlas.claim_index import CONFIG, index, load  # noqa: E402


def test_config_is_tracked_and_well_formed():
    assert CONFIG.exists(), "config/claim_restatements.yaml must be committed"
    specs = load()
    assert specs, "no manuscripts listed"
    for name, spec in specs.items():
        assert spec.get("file", "").endswith(".docx"), name
        assert spec.get("claims"), f"{name} lists no claims"
        ids = [c["id"] for c in spec["claims"]]
        assert len(ids) == len(set(ids)), f"{name} has duplicate claim ids: {ids}"
        for c in spec["claims"]:
            assert c.get("says", "").strip(), f"{name}/{c['id']} has no plain-language statement"
            assert c.get("pattern", "").strip(), f"{name}/{c['id']} has no pattern"


def test_every_pattern_compiles():
    for name, spec in load().items():
        for c in spec["claims"]:
            try:
                re.compile(c["pattern"], re.I)
            except re.error as exc:
                pytest.fail(f"{name}/{c['id']}: {exc}")


def test_both_manuscripts_are_covered():
    """Both papers draw on the same genes and the same method family, so a claim
    corrected in one is worth checking in the other."""
    files = {s["file"] for s in load().values()}
    assert len(files) >= 2, f"only {files} tracked; both manuscripts should be"


def test_index_finds_multiple_restatements_when_the_manuscript_is_available():
    """The failure this exists to prevent is a claim restated somewhere nobody
    looked, so a claim that resolves to a single location is the suspicious
    case, not the reassuring one."""
    specs = load()
    for name, spec in specs.items():
        path = Path.home() / "Desktop" / spec["file"]
        if not path.exists():
            continue
        idx = index(path, spec["claims"])
        assert any(len(v) > 1 for v in idx.values()), (
            f"{name}: no claim resolves to more than one location, which means the "
            "patterns are too narrow to catch a restatement")
        return
    pytest.skip("neither manuscript is available in this checkout")


def _blocks(path: Path) -> list[str]:
    """Paragraph and table-cell text, the same surface claim_index searches."""
    from docx import Document

    d = Document(str(path))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.extend(c.text for c in row.cells)
    return out


def _near_miss_variants(term: str) -> set[str]:
    """Spelling variants a pattern written one way would silently miss."""
    v = {term}
    if "-" in term:
        v |= {term.replace("-", " "), term.replace("-", "")}
    if " " in term:
        v.add(term.replace(" ", "-"))
    v.add(term[:-1] if term.endswith("s") else term + "s")
    return v - {term}


def test_patterns_have_no_near_miss_spelling_variants():
    """A pattern that matches 'deep intron' but not 'deep-intronic' reports a
    claim as stated in nine places when it is stated in fifteen, and a
    correction applied to it misses the six. The failure is silent, so it is
    checked rather than watched for: for each literal alternative in each
    pattern, a hyphen, spacing or plural variant that appears in the manuscript
    and is NOT matched by the pattern is a gap in the index.
    """
    import re

    specs = load()
    checked = 0
    gaps = []
    for name, spec in specs.items():
        path = Path.home() / "Desktop" / spec["file"]
        if not path.exists():
            continue
        checked += 1
        blocks = _blocks(path)
        for claim in spec["claims"]:
            pat = re.compile(claim["pattern"], re.I)
            for alt in claim["pattern"].split("|"):
                literal = re.sub(r"[()\[\]?*+^$\\]|\{\d+,?\d*\}", "", alt).strip()
                if len(literal) < 5:
                    continue
                for variant in _near_miss_variants(literal):
                    hit = next((t for t in blocks
                                if variant.lower() in t.lower() and not pat.search(t)), None)
                    if hit:
                        gaps.append(f"{name}/{claim['id']}: pattern has {literal!r} but the "
                                    f"manuscript also says {variant!r} — e.g. {hit[:70]!r}")
    if not checked:
        pytest.skip("neither manuscript is available in this checkout")
    assert not gaps, "restatement patterns with spelling gaps:\n  " + "\n  ".join(gaps)
